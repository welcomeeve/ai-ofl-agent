import os
import io
import threading
import logging
import requests
from flask import Flask, request, send_file, render_template_string
from docx import Document
from telegram import Update
from telegram.ext import Application, CommandHandler, MessageHandler, ContextTypes, filters

logging.basicConfig(level=logging.INFO)
app = Flask(__name__)

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
ALLOWED_TELEGRAM_USER_ID = os.getenv("ALLOWED_TELEGRAM_USER_ID", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash").strip()

HTML = """
<!doctype html>
<html lang="bn">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>AI Language & Document Agent</title>
<style>
body{font-family:system-ui,Arial,sans-serif;max-width:760px;margin:24px auto;padding:0 16px;line-height:1.5}
input,select,textarea,button{box-sizing:border-box;width:100%;padding:11px;margin:6px 0 14px;font:inherit}
textarea{min-height:110px}button{cursor:pointer;background:#202938;color:white;border:0;border-radius:8px}
.small{color:#555;font-size:.92rem}
</style>
</head>
<body>
<h1>AI Language &amp; Document Agent</h1>
<p>টার্গেট ভাষা নির্ধারণ করে অনুবাদ বা Word ডকুমেন্ট তৈরি করুন।</p>
<form method="post" action="/generate">
<label>কাজের ধরন</label>
<select name="document_type">
<option>Offer Letter</option><option>Cover Letter</option>
<option>Business Letter</option><option>Translation</option><option>Other Professional Document</option>
</select>
<label>টার্গেট ভাষা (যেমন French, Arabic, German)</label>
<input name="target_language" required placeholder="French">
<label>প্রাপকের নাম (ঐচ্ছিক)</label>
<input name="person_name" placeholder="Marie Dupont">
<label>প্রাপকের gender/সম্বোধন (ঐচ্ছিক; না জানলে Unknown)</label>
<select name="gender">
<option>Unknown / not specified</option><option>Woman</option><option>Man</option><option>Non-binary / neutral wording preferred</option>
</select>
<label>কোম্পানির নাম (ঐচ্ছিক)</label>
<input name="company" placeholder="ABC Solutions">
<label>পদবি / চাকরির পদ (ঐচ্ছিক)</label>
<input name="job_title" placeholder="Operations Manager">
<label>তারিখ, বেতন, শর্তাবলি এবং অন্যান্য তথ্য</label>
<textarea name="details" required placeholder="যে তথ্যগুলো চিঠিতে থাকতে হবে লিখুন। অজানা তথ্য বানিয়ে দিতে বলবেন না।"></textarea>
<label>লেখার ধরন</label>
<select name="tone"><option>Formal and professional</option><option>Warm and professional</option><option>Concise</option></select>
<button type="submit">Word (.docx) তৈরি করুন</button>
</form>
<p class="small">গুরুত্বপূর্ণ: অফার লেটার তৈরির আগে কোম্পানির অনুমোদন এবং তথ্য যাচাই করুন। API key কখনো ওয়েব ফর্মে লিখবেন না।</p>
</body></html>
"""

def make_prompt(data):
    return f"""
You are a careful multilingual translator and professional document drafter.
Create the requested document in the target language. Follow the target language's
grammar, spelling, honorifics, grammatical gender, agreement, and job-title conventions.
If the recipient's gender is unknown, prefer natural gender-neutral wording where possible
and do not infer gender from a name alone. Preserve names, dates, amounts, and facts exactly.
Do not invent salary, legal clauses, signatures, addresses, or company facts. If key details
are missing, use clearly marked [DETAIL REQUIRED] placeholders. Do not claim that a job offer
is genuine or approved. Return only the finished document text, with sensible headings and
paragraphs, no commentary about the process.

Document type: {data.get('document_type','Offer Letter')}
Target language: {data.get('target_language','')}
Recipient name: {data.get('person_name','') or 'Not supplied'}
Recipient gender/preferred wording: {data.get('gender','Unknown / not specified')}
Company: {data.get('company','') or 'Not supplied'}
Job title: {data.get('job_title','') or 'Not supplied'}
Tone: {data.get('tone','Formal and professional')}
Client instructions and facts:
{data.get('details','')}
""".strip()

def generate_text(data):
    if not GEMINI_API_KEY:
        raise RuntimeError("GEMINI_API_KEY সেট করা নেই। Hosting-এর Environment Variables-এ API key যোগ করুন।")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent"
    response = requests.post(
        url,
        params={"key": GEMINI_API_KEY},
        json={
            "contents": [{"parts": [{"text": make_prompt(data)}]}],
            "generationConfig": {"temperature": 0.3, "maxOutputTokens": 5000}
        },
        timeout=90
    )
    if response.status_code != 200:
        logging.error("Gemini API error %s: %s", response.status_code, response.text[:1000])
        raise RuntimeError(f"Gemini API-তে সমস্যা হয়েছে (HTTP {response.status_code})। API key, model ও quota পরীক্ষা করুন।")
    payload = response.json()
    try:
        return "".join(part.get("text", "") for part in payload["candidates"][0]["content"]["parts"]).strip()
    except (KeyError, IndexError, TypeError):
        raise RuntimeError("Gemini থেকে প্রত্যাশিত লেখা পাওয়া যায়নি। আবার চেষ্টা করুন।")

def create_docx(text, title="AI Document"):
    doc = Document()
    doc.add_heading(title, level=1)
    for paragraph in text.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream

@app.get("/")
def home():
    return render_template_string(HTML)

@app.post("/generate")
def web_generate():
    data = {
        "document_type": request.form.get("document_type", "Offer Letter"),
        "target_language": request.form.get("target_language", "").strip(),
        "person_name": request.form.get("person_name", "").strip(),
        "gender": request.form.get("gender", "Unknown / not specified"),
        "company": request.form.get("company", "").strip(),
        "job_title": request.form.get("job_title", "").strip(),
        "details": request.form.get("details", "").strip(),
        "tone": request.form.get("tone", "Formal and professional"),
    }
    if not data["target_language"] or not data["details"]:
        return "টার্গেট ভাষা ও নির্দেশনা দিতে হবে।", 400
    try:
        output = generate_text(data)
        docx = create_docx(output, data["document_type"])
        filename = (data["document_type"].lower().replace(" ", "_") or "document") + ".docx"
        return send_file(docx, as_attachment=True, download_name=filename,
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as exc:
        logging.exception("Document generation failed")
        return f"ডকুমেন্ট তৈরি হয়নি: {str(exc)}", 500

def parse_telegram_command(text):
    # User sends a compact form after /make:
    # /make
    # Type: Offer Letter
    # Language: French
    # Name: Marie Dupont
    # Gender: Woman
    # Company: ABC Solutions
    # Job title: Operations Manager
    # Tone: Formal and professional
    # Details: ...
    data = {
        "document_type": "Offer Letter",
        "target_language": "",
        "person_name": "",
        "gender": "Unknown / not specified",
        "company": "",
        "job_title": "",
        "tone": "Formal and professional",
        "details": ""
    }
    current_key = None
    details_lines = []
    keys = {
        "type": "document_type", "language": "target_language", "name": "person_name",
        "gender": "gender", "company": "company", "job title": "job_title", "tone": "tone"
    }
    for line in text.splitlines()[1:]:
        if ":" in line:
            label, value = line.split(":", 1)
            normalized = label.strip().lower()
            if normalized in keys:
                current_key = keys[normalized]
                data[current_key] = value.strip()
                continue
            if normalized == "details":
                current_key = "details"
                details_lines.append(value.strip())
                continue
        if current_key == "details":
            details_lines.append(line)
    data["details"] = "\n".join(x for x in details_lines if x).strip()
    return data

async def start_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "স্বাগতম! Word ডকুমেন্ট বানাতে /make লিখে নিচের ফরম্যাট পূরণ করে পাঠান:\n\n"
        "/make\nType: Offer Letter\nLanguage: French\nName: Marie Dupont\n"
        "Gender: Woman\nCompany: ABC Solutions\nJob title: Operations Manager\n"
        "Tone: Formal and professional\nDetails: Start date is 1 December 2026. Include only these facts.\n\n"
        "ধরন: Offer Letter, Cover Letter, Business Letter, Translation"
    )

async def make_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    if ALLOWED_TELEGRAM_USER_ID and user_id != ALLOWED_TELEGRAM_USER_ID:
        await update.message.reply_text("এই বটটি ব্যবহার করার অনুমতি আপনার নেই।")
        return
    raw = update.message.text or ""
    data = parse_telegram_command(raw)
    if not data["target_language"] or not data["details"]:
        await update.message.reply_text("ভাষা ও Details দিতে হবে। /start লিখে উদাহরণ দেখুন।")
        return
    await update.message.reply_text("ডকুমেন্ট তৈরি করছি—একটু অপেক্ষা করুন।")
    try:
        text = generate_text(data)
        stream = create_docx(text, data["document_type"])
        filename = (data["document_type"].lower().replace(" ", "_") or "document") + ".docx"
        await update.message.reply_document(document=stream, filename=filename, caption="আপনার Word ডকুমেন্ট তৈরি হয়েছে। তথ্য যাচাই করে ব্যবহার করুন।")
    except Exception:
        logging.exception("Telegram generation failed")
        await update.message.reply_text("ডকুমেন্ট তৈরি করা যায়নি। API key, model, quota এবং hosting logs পরীক্ষা করুন।")

async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("ব্যবহার: /start লিখুন। তারপর /make কমান্ডের ফরম্যাট অনুসরণ করুন।")

def run_telegram_bot():
    if not TELEGRAM_BOT_TOKEN:
        logging.warning("TELEGRAM_BOT_TOKEN নেই; Telegram bot চালু হবে না। Website চলতে পারে।")
        return
    async def runner():
        application = Application.builder().token(TELEGRAM_BOT_TOKEN).build()
        application.add_handler(CommandHandler("start", start_cmd))
        application.add_handler(CommandHandler("help", help_cmd))
        application.add_handler(MessageHandler(filters.COMMAND & filters.Regex(r"^/make"), make_cmd))
        await application.initialize()
        await application.start()
        await application.updater.start_polling(drop_pending_updates=True)
        # Keep the polling task alive for the lifetime of the process.
        import asyncio
        while True:
            await asyncio.sleep(3600)
    import asyncio
    asyncio.run(runner())

# Gunicorn imports this module, so start the bot on import when configured.
# Keep the hosting service to one web worker to avoid duplicate polling instances.
if TELEGRAM_BOT_TOKEN and os.getenv("RUN_TELEGRAM_BOT", "1") == "1":
    threading.Thread(target=run_telegram_bot, daemon=True, name="telegram-bot").start()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "8080")))
